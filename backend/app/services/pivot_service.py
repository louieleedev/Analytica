from __future__ import annotations

from io import BytesIO
import logging
import re
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill

from app.db.duckdb import get_connection
from app.services.dataset_profile import _numeric_expression
from app.services.dataset_storage import get_dataset_storage_info, quote_identifier
from app.services.explorer_profile import _build_where_clause


logger = logging.getLogger(__name__)

AMOUNT_AGGREGATIONS = {"Sum", "Average", "Median", "Min", "Max", "Count", "Distinct Count"}
DISCRETE_AGGREGATIONS = {"Count", "Distinct Count"}
RESULT_LIMIT = 10_000
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
HEADER_FILL = "E5E7EB"
TOTAL_FILL = "E5E7EB"


def estimate_pivot(dataset_id: str, payload: dict[str, Any]) -> dict[str, Any]:
    storage_info = get_dataset_storage_info(dataset_id)
    if storage_info is None:
        raise HTTPException(status_code=404, detail=f"Dataset {dataset_id} was not found.")

    row_fields = _normalise_field_list(payload.get("rows", []))
    column_fields = _normalise_field_list(payload.get("columns", []))
    value_fields = _normalise_field_list(payload.get("values", []))
    filters = payload.get("filters", [])

    _validate_columns(storage_info, [field["name"] for field in row_fields])
    _validate_columns(storage_info, [field["name"] for field in column_fields])
    _validate_columns(storage_info, [field["name"] for field in value_fields])
    _validate_columns(
        storage_info,
        [str(filter_payload.get("columnName")) for filter_payload in filters if filter_payload.get("columnName")],
    )

    where_sql, parameters = _build_where_clause(storage_info, filters)
    row_cardinalities = _estimate_field_cardinalities(storage_info, [field["name"] for field in row_fields], where_sql, parameters)
    column_cardinalities = _estimate_field_cardinalities(
        storage_info,
        [field["name"] for field in column_fields],
        where_sql,
        parameters,
    )
    row_combinations = _multiply_cardinalities(row_cardinalities)
    column_combinations = _multiply_cardinalities(column_cardinalities)
    value_count = max(1, len(value_fields))
    projected_columns = (
        len(row_fields) + max(1, column_combinations) * value_count + value_count
        if column_fields
        else len(row_fields) + value_count
    )

    logger.info(
        "Pivot estimate generated: dataset_id=%s row_combinations=%s column_combinations=%s projected_columns=%s",
        dataset_id,
        row_combinations,
        column_combinations,
        projected_columns,
    )

    return {
        "datasetId": dataset_id,
        "rowCombinations": row_combinations,
        "columnCombinations": column_combinations,
        "projectedColumns": projected_columns,
        "valueCount": len(value_fields),
        "rowCardinalities": row_cardinalities,
        "columnCardinalities": column_cardinalities,
    }


def execute_pivot(dataset_id: str, payload: dict[str, Any]) -> dict[str, Any]:
    return _execute_pivot(dataset_id, payload)


def build_pivot_export(payload: dict[str, Any]) -> tuple[bytes, str]:
    dataset_id = str(payload.get("datasetId") or "").strip()
    if not dataset_id:
        raise HTTPException(status_code=400, detail="datasetId is required.")

    project_name = str(payload.get("projectName") or "Project").strip() or "Project"
    pivot_payload = {
        "rows": payload.get("rows", []),
        "columns": payload.get("columns", []),
        "values": payload.get("values", []),
        "filters": payload.get("filters", []),
    }
    pivot_result = _execute_pivot(dataset_id, pivot_payload)
    _validate_excel_limits(pivot_result)

    workbook = _build_pivot_workbook(pivot_result)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream.getvalue(), _build_export_filename(project_name)


def _execute_pivot(dataset_id: str, payload: dict[str, Any]) -> dict[str, Any]:
    storage_info = get_dataset_storage_info(dataset_id)
    if storage_info is None:
        raise HTTPException(status_code=404, detail=f"Dataset {dataset_id} was not found.")

    row_fields = _normalise_field_list(payload.get("rows", []))
    value_fields = _normalise_field_list(payload.get("values", []))
    filters = payload.get("filters", [])
    columns = _normalise_field_list(payload.get("columns", []))

    if not value_fields:
        raise HTTPException(status_code=400, detail="At least one value field is required.")
    if not row_fields:
        raise HTTPException(status_code=400, detail="At least one row field is required.")

    _validate_columns(storage_info, [field["name"] for field in row_fields])
    _validate_columns(storage_info, [field["name"] for field in value_fields])
    _validate_columns(storage_info, [field["name"] for field in columns])
    _validate_columns(
        storage_info,
        [str(filter_payload.get("columnName")) for filter_payload in filters if filter_payload.get("columnName")],
    )

    if columns:
        return _execute_column_pivot(dataset_id, storage_info, row_fields, columns, value_fields, filters)

    select_parts = [quote_identifier(field["name"]) for field in row_fields]
    group_by_parts = [quote_identifier(field["name"]) for field in row_fields]
    value_headers: list[str] = []

    for value_field in value_fields:
        aggregation = str(value_field.get("aggregation") or "Count")
        category = str(value_field.get("category") or "")
        column_name = str(value_field["name"])
        _validate_aggregation(category, aggregation, column_name)
        header = _unique_header(f"{column_name} {aggregation}", [field["name"] for field in row_fields] + value_headers)
        value_headers.append(header)
        select_parts.append(f"{_aggregation_expression(column_name, category, aggregation)} AS {quote_identifier(header)}")

    where_sql, parameters = _build_where_clause(storage_info, filters)
    query = _build_query(storage_info, select_parts, group_by_parts, where_sql, limit=None)

    logger.info(
        "Pivot query generated: dataset_id=%s rows=%s values=%s sql=%s",
        dataset_id,
        [field["name"] for field in row_fields],
        value_headers,
        " ".join(query.split()),
    )

    for connection in get_connection():
        result = connection.execute(query, parameters).fetchall()

    headers = [field["name"] for field in row_fields] + value_headers
    rows = [[_serialise_value(value) for value in row] for row in result]
    return {
        "datasetId": dataset_id,
        "headers": headers,
        "rows": rows,
        "rowCount": len(rows),
        "warnings": [],
        "sql": query,
    }


def _execute_column_pivot(
    dataset_id: str,
    storage_info: dict[str, Any],
    row_fields: list[dict[str, Any]],
    column_fields: list[dict[str, Any]],
    value_fields: list[dict[str, Any]],
    filters: list[dict[str, Any]],
) -> dict[str, Any]:
    column_field_names = [str(field["name"]) for field in column_fields]
    row_field_names = [str(field["name"]) for field in row_fields]
    where_sql, parameters = _build_where_clause(storage_info, filters)

    value_headers: list[str] = []
    metric_select_parts: list[str] = []
    for value_field in value_fields:
        aggregation = str(value_field.get("aggregation") or "Count")
        category = str(value_field.get("category") or "")
        metric_column = str(value_field["name"])
        _validate_aggregation(category, aggregation, metric_column)
        metric_header = _unique_header(f"{metric_column} {aggregation}", value_headers)
        value_headers.append(metric_header)
        metric_select_parts.append(
            f"{_aggregation_expression(metric_column, category, aggregation)} AS {quote_identifier(metric_header)}"
        )

    column_values = _query_distinct_column_values(
        storage_info,
        column_field_names,
        value_fields,
        where_sql,
        parameters,
        limit=None,
    )

    select_parts = [
        *[quote_identifier(field_name) for field_name in row_field_names],
        *[quote_identifier(field_name) for field_name in column_field_names],
        *metric_select_parts,
    ]
    group_by_parts = [
        *[quote_identifier(field_name) for field_name in row_field_names],
        *[quote_identifier(field_name) for field_name in column_field_names],
    ]
    query = _build_query(storage_info, select_parts, group_by_parts, where_sql, limit=None)

    logger.info(
        "Column pivot query generated: dataset_id=%s rows=%s columns=%s values=%s sql=%s",
        dataset_id,
        row_field_names,
        column_field_names,
        value_headers,
        " ".join(query.split()),
    )

    for connection in get_connection():
        result = connection.execute(query, parameters).fetchall()

    headers = _build_column_pivot_headers(row_field_names, column_values, value_headers)
    row_totals = _query_row_totals(storage_info, row_field_names, metric_select_parts, where_sql, parameters)
    column_totals = _query_column_totals(storage_info, column_field_names, metric_select_parts, where_sql, parameters)
    grand_totals = _query_grand_totals(storage_info, metric_select_parts, where_sql, parameters)
    rows = _build_column_pivot_rows(
        result,
        row_field_count=len(row_field_names),
        column_field_count=len(column_field_names),
        column_values=column_values,
        value_headers=value_headers,
        row_totals=row_totals,
    )
    total_row = _build_column_total_row(
        row_field_count=len(row_field_names),
        column_values=column_values,
        value_headers=value_headers,
        column_totals=column_totals,
        grand_totals=grand_totals,
    )
    if total_row:
        rows.append(total_row)
    return {
        "datasetId": dataset_id,
        "headers": headers,
        "rows": rows,
        "rowCount": len(rows),
        "warnings": [],
        "sql": query,
        "pivot": {
            "rowFields": row_field_names,
            "columnFields": column_field_names,
            "columnValues": [[_serialise_value(value) for value in values] for values in column_values],
            "valueFields": value_headers,
        },
    }


def _normalise_field_list(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []

    fields = []
    for item in value:
        if isinstance(item, str):
            fields.append({"name": item})
        elif isinstance(item, dict) and item.get("name"):
            fields.append(item)
    return fields


def _validate_columns(storage_info: dict[str, Any], column_names: list[str]) -> None:
    available_columns = {str(column["name"]) for column in storage_info["columns"]}
    missing_columns = [column_name for column_name in column_names if column_name not in available_columns]
    if missing_columns:
        raise HTTPException(status_code=404, detail=f"Columns were not found: {', '.join(missing_columns)}")


def _validate_aggregation(category: str, aggregation: str, column_name: str) -> None:
    if category == "AMOUNT":
        allowed = AMOUNT_AGGREGATIONS
    else:
        allowed = DISCRETE_AGGREGATIONS

    if aggregation not in allowed:
        raise HTTPException(
            status_code=400,
            detail=f"{aggregation} is not supported for {column_name}.",
        )


def _aggregation_expression(column_name: str, category: str, aggregation: str) -> str:
    quoted_column = quote_identifier(column_name)
    value_expression = _numeric_expression(column_name) if category == "AMOUNT" else quoted_column

    if aggregation == "Sum":
        return f"SUM({value_expression})"
    if aggregation == "Average":
        return f"AVG({value_expression})"
    if aggregation == "Median":
        return f"MEDIAN({value_expression})"
    if aggregation == "Min":
        return f"MIN({value_expression})"
    if aggregation == "Max":
        return f"MAX({value_expression})"
    if aggregation == "Distinct Count":
        return f"COUNT(DISTINCT {quoted_column})"

    return f"COUNT({quoted_column})"


def _unique_header(base_header: str, existing_headers: list[str]) -> str:
    if base_header not in existing_headers:
        return base_header

    suffix = 2
    while f"{base_header} {suffix}" in existing_headers:
        suffix += 1

    return f"{base_header} {suffix}"


def _build_query(
    storage_info: dict[str, Any],
    select_parts: list[str],
    group_by_parts: list[str],
    where_sql: str,
    limit: int | None = RESULT_LIMIT,
) -> str:
    group_by_sql = f"GROUP BY {', '.join(group_by_parts)}" if group_by_parts else ""
    order_by_sql = f"ORDER BY {', '.join(group_by_parts)}" if group_by_parts else ""
    limit_sql = f"LIMIT {limit}" if limit is not None else ""

    return f"""
        SELECT {', '.join(select_parts)}
        FROM {quote_identifier(str(storage_info["tableName"]))}
        {where_sql}
        {group_by_sql}
        {order_by_sql}
        {limit_sql}
    """


def _query_distinct_column_values(
    storage_info: dict[str, Any],
    column_names: list[str],
    value_fields: list[dict[str, Any]],
    where_sql: str,
    parameters: list[Any],
    limit: int | None = RESULT_LIMIT,
) -> list[tuple[Any, ...]]:
    quoted_columns = [quote_identifier(column_name) for column_name in column_names]
    selected_columns = ", ".join(quoted_columns)
    value_presence_sql = _value_presence_condition(value_fields)
    distinct_where_sql = _append_where_condition(where_sql, value_presence_sql)
    limit_sql = "LIMIT ?" if limit is not None else ""
    query_parameters = [*parameters, limit] if limit is not None else parameters
    for connection in get_connection():
        rows = connection.execute(
            f"""
            SELECT DISTINCT {selected_columns}
            FROM {quote_identifier(str(storage_info["tableName"]))}
            {distinct_where_sql}
            ORDER BY {selected_columns}
            {limit_sql}
            """,
            query_parameters,
        ).fetchall()
        return [tuple(row) for row in rows]

    return []


def _value_presence_condition(value_fields: list[dict[str, Any]]) -> str:
    conditions = []
    for value_field in value_fields:
        column_name = str(value_field["name"])
        category = str(value_field.get("category") or "")
        expression = _numeric_expression(column_name) if category == "AMOUNT" else quote_identifier(column_name)
        conditions.append(f"{expression} IS NOT NULL")

    return " OR ".join(f"({condition})" for condition in conditions)


def _append_where_condition(where_sql: str, condition_sql: str) -> str:
    if not condition_sql:
        return where_sql
    if where_sql:
        return f"{where_sql} AND ({condition_sql})"
    return f"WHERE ({condition_sql})"


def _estimate_field_cardinalities(
    storage_info: dict[str, Any],
    column_names: list[str],
    where_sql: str,
    parameters: list[Any],
) -> dict[str, int]:
    cardinalities: dict[str, int] = {}
    table_name = quote_identifier(str(storage_info["tableName"]))

    for column_name in column_names:
        quoted_column = quote_identifier(column_name)
        for connection in get_connection():
            row = connection.execute(
                f"""
                SELECT
                    COUNT(DISTINCT {quoted_column})
                    + CASE WHEN COUNT(*) FILTER (WHERE {quoted_column} IS NULL) > 0 THEN 1 ELSE 0 END
                FROM {table_name}
                {where_sql}
                """,
                parameters,
            ).fetchone()
            cardinalities[column_name] = int(row[0] or 0) if row else 0

    return cardinalities


def _multiply_cardinalities(cardinalities: dict[str, int]) -> int:
    if not cardinalities:
        return 0

    product = 1
    for value in cardinalities.values():
        product *= max(0, int(value))
    return product


def _is_total_row(row: list[Any]) -> bool:
    if not row:
        return False

    label = str(row[0])
    return label == "Total" or label.startswith("Total ") or label.startswith("Grand ")


def _build_column_pivot_headers(
    row_field_names: list[str],
    column_values: list[tuple[Any, ...]],
    value_headers: list[str],
) -> list[str]:
    headers = [*row_field_names]
    multiple_values = len(value_headers) > 1

    for column_value in column_values:
        column_label = _column_tuple_label(column_value)
        for value_header in value_headers:
            metric_label = _metric_label(value_header)
            base_header = f"{column_label} {metric_label}" if multiple_values else column_label
            header = _unique_header(base_header, headers)
            headers.append(header)

    for value_header in value_headers:
        metric_label = _metric_label(value_header)
        base_header = _total_header_label(metric_label, multiple_values)
        headers.append(_unique_header(base_header, headers))

    return headers


def _build_column_pivot_rows(
    result: list[tuple[Any, ...]],
    row_field_count: int,
    column_field_count: int,
    column_values: list[tuple[Any, ...]],
    value_headers: list[str],
    row_totals: dict[tuple[Any, ...], list[Any]],
) -> list[list[Any]]:
    row_map: dict[tuple[Any, ...], dict[tuple[Any, ...], list[Any]]] = {}
    row_order: list[tuple[Any, ...]] = []

    for result_row in result:
        row_key = tuple(result_row[:row_field_count])
        column_value = tuple(result_row[row_field_count : row_field_count + column_field_count])
        metric_values = list(result_row[row_field_count + column_field_count :])

        if row_key not in row_map:
            row_map[row_key] = {}
            row_order.append(row_key)
        row_map[row_key][column_value] = metric_values

    rows: list[list[Any]] = []
    for row_key in row_order:
        output_row = [_serialise_value(value) for value in row_key]
        for column_value in column_values:
            metric_values = row_map[row_key].get(column_value)
            if metric_values is None:
                output_row.extend(0 for _ in value_headers)
            else:
                output_row.extend(_serialise_value(value) if value is not None else 0 for value in metric_values)
        output_row.extend(
            _serialise_value(value) if value is not None else 0
            for value in row_totals.get(row_key, [0 for _ in value_headers])
        )
        rows.append(output_row)

    return rows


def _build_column_total_row(
    row_field_count: int,
    column_values: list[tuple[Any, ...]],
    value_headers: list[str],
    column_totals: dict[tuple[Any, ...], list[Any]],
    grand_totals: list[Any],
) -> list[Any] | None:
    if not column_values and not grand_totals:
        return None

    row: list[Any] = [_total_row_label(value_headers)]
    row.extend("" for _ in range(max(0, row_field_count - 1)))
    for column_value in column_values:
        row.extend(
            _serialise_value(value) if value is not None else 0
            for value in column_totals.get(column_value, [0 for _ in value_headers])
        )
    row.extend(_serialise_value(value) if value is not None else 0 for value in grand_totals)
    return row


def _query_row_totals(
    storage_info: dict[str, Any],
    row_field_names: list[str],
    metric_select_parts: list[str],
    where_sql: str,
    parameters: list[Any],
) -> dict[tuple[Any, ...], list[Any]]:
    select_parts = [*[quote_identifier(field_name) for field_name in row_field_names], *metric_select_parts]
    group_by_parts = [quote_identifier(field_name) for field_name in row_field_names]
    query = _build_query(storage_info, select_parts, group_by_parts, where_sql, limit=None)

    for connection in get_connection():
        rows = connection.execute(query, parameters).fetchall()
        return {
            tuple(row[: len(row_field_names)]): list(row[len(row_field_names) :])
            for row in rows
        }

    return {}


def _query_column_totals(
    storage_info: dict[str, Any],
    column_names: list[str],
    metric_select_parts: list[str],
    where_sql: str,
    parameters: list[Any],
) -> dict[tuple[Any, ...], list[Any]]:
    select_parts = [*[quote_identifier(column_name) for column_name in column_names], *metric_select_parts]
    group_by_parts = [quote_identifier(column_name) for column_name in column_names]
    query = _build_query(storage_info, select_parts, group_by_parts, where_sql, limit=None)

    for connection in get_connection():
        rows = connection.execute(query, parameters).fetchall()
        return {tuple(row[: len(column_names)]): list(row[len(column_names) :]) for row in rows}

    return {}


def _query_grand_totals(
    storage_info: dict[str, Any],
    metric_select_parts: list[str],
    where_sql: str,
    parameters: list[Any],
) -> list[Any]:
    query = _build_query(storage_info, metric_select_parts, [], where_sql, limit=None)

    for connection in get_connection():
        row = connection.execute(query, parameters).fetchone()
        return [] if row is None else list(row)

    return []


def _metric_label(value_header: str) -> str:
    for aggregation in sorted(AMOUNT_AGGREGATIONS | DISCRETE_AGGREGATIONS, key=len, reverse=True):
        suffix = f" {aggregation}"
        if value_header.endswith(suffix):
            return aggregation
    return value_header


def _column_tuple_label(column_value: tuple[Any, ...]) -> str:
    return " ".join(_column_value_label(value) for value in column_value)


def _column_value_label(value: Any) -> str:
    serialised_value = _serialise_value(value)
    if serialised_value is None:
        return "(Blank)"
    text = str(serialised_value).strip()
    return text if text else "(Blank)"


def _total_header_label(metric_label: str, multiple_values: bool) -> str:
    if metric_label in {"Sum", "Count"}:
        return f"Total {metric_label}" if multiple_values else "Total"
    return f"Grand {metric_label}"


def _total_row_label(value_headers: list[str]) -> str:
    metric_labels = [_metric_label(value_header) for value_header in value_headers]
    if len(metric_labels) == 1:
        return _total_header_label(metric_labels[0], multiple_values=False)
    if all(metric_label in {"Sum", "Count"} for metric_label in metric_labels):
        return "Total"
    return "Grand Total"


def _validate_excel_limits(pivot_result: dict[str, Any]) -> None:
    header_depth = _excel_header_depth(pivot_result)
    worksheet_rows = header_depth + len(pivot_result["rows"])
    worksheet_columns = len(pivot_result["headers"])
    if worksheet_rows > EXCEL_MAX_ROWS or worksheet_columns > EXCEL_MAX_COLUMNS:
        raise HTTPException(status_code=400, detail="Pivot result exceeds Excel worksheet limits.")


def _build_pivot_workbook(pivot_result: dict[str, Any]) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Pivot"

    header_rows = _excel_header_rows(pivot_result)
    for row_index, row in enumerate(header_rows, start=1):
        for column_index, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            _style_header_cell(cell)

    _merge_excel_headers(worksheet, pivot_result, header_rows)

    body_start_row = len(header_rows) + 1
    for row_offset, row in enumerate(pivot_result["rows"], start=0):
        excel_row = body_start_row + row_offset
        is_total_row = _is_total_row(row)
        for column_index, value in enumerate(row, start=1):
            cell = worksheet.cell(row=excel_row, column=column_index, value=value)
            if is_total_row or _is_total_header(pivot_result["headers"][column_index - 1]):
                _style_total_cell(cell)

    for column_index, column_cells in enumerate(worksheet.columns, start=1):
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:100])
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 32)

    worksheet.freeze_panes = worksheet.cell(row=body_start_row, column=2)
    return workbook


def _excel_header_depth(pivot_result: dict[str, Any]) -> int:
    pivot_metadata = pivot_result.get("pivot")
    if not pivot_metadata or not pivot_metadata.get("columnFields") or not pivot_metadata.get("columnValues"):
        return 1

    value_headers = pivot_metadata.get("valueFields") or ["Value"]
    return len(pivot_metadata["columnFields"]) + (1 if len(value_headers) > 1 else 0)


def _excel_header_rows(pivot_result: dict[str, Any]) -> list[list[Any]]:
    header_depth = _excel_header_depth(pivot_result)
    headers = pivot_result["headers"]
    pivot_metadata = pivot_result.get("pivot")
    if not pivot_metadata or not pivot_metadata.get("columnFields") or not pivot_metadata.get("columnValues"):
        return [headers]

    row_fields = pivot_metadata["rowFields"]
    column_values = pivot_metadata["columnValues"]
    value_headers = pivot_metadata.get("valueFields") or ["Value"]
    value_count = len(value_headers)
    show_metric_level = value_count > 1
    rows = [["" for _ in headers] for _ in range(header_depth)]

    for index, row_field in enumerate(row_fields):
        rows[0][index] = row_field

    for level in range(len(pivot_metadata["columnFields"])):
        for combination_index, column_value in enumerate(column_values):
            rows[level][len(row_fields) + combination_index * value_count] = _excel_header_value(column_value[level])

    if show_metric_level:
        metric_row = header_depth - 1
        for combination_index, _ in enumerate(column_values):
            for value_index, value_header in enumerate(value_headers):
                rows[metric_row][len(row_fields) + combination_index * value_count + value_index] = _metric_label(value_header)

    total_start = len(row_fields) + len(column_values) * value_count
    total_headers = headers[total_start:]
    if total_headers:
        rows[0][total_start] = "Total" if show_metric_level else total_headers[0]
        if show_metric_level:
            for index, total_header in enumerate(total_headers):
                rows[-1][total_start + index] = _metric_label(total_header)

    return rows


def _merge_excel_headers(worksheet: Any, pivot_result: dict[str, Any], header_rows: list[list[Any]]) -> None:
    pivot_metadata = pivot_result.get("pivot")
    if not pivot_metadata or not pivot_metadata.get("columnFields") or not pivot_metadata.get("columnValues"):
        return

    row_field_count = len(pivot_metadata["rowFields"])
    column_values = pivot_metadata["columnValues"]
    value_headers = pivot_metadata.get("valueFields") or ["Value"]
    value_count = len(value_headers)
    header_depth = len(header_rows)
    show_metric_level = value_count > 1

    for index in range(row_field_count):
        if header_depth > 1:
            worksheet.merge_cells(start_row=1, start_column=index + 1, end_row=header_depth, end_column=index + 1)

    for level in range(len(pivot_metadata["columnFields"])):
        group_start = 0
        while group_start < len(column_values):
            group_end = group_start + 1
            while group_end < len(column_values) and _same_header_prefix(column_values[group_start], column_values[group_end], level):
                group_end += 1

            start_column = row_field_count + group_start * value_count + 1
            end_column = row_field_count + group_end * value_count
            if end_column > start_column:
                worksheet.merge_cells(start_row=level + 1, start_column=start_column, end_row=level + 1, end_column=end_column)
            group_start = group_end

    total_start = row_field_count + len(column_values) * value_count + 1
    total_headers = pivot_result["headers"][total_start - 1 :]
    if total_headers:
        if show_metric_level and header_depth > 1:
            end_column = total_start + len(total_headers) - 1
            if end_column > total_start:
                worksheet.merge_cells(start_row=1, start_column=total_start, end_row=header_depth - 1, end_column=end_column)
            else:
                worksheet.merge_cells(start_row=1, start_column=total_start, end_row=header_depth - 1, end_column=total_start)
        elif header_depth > 1:
            worksheet.merge_cells(start_row=1, start_column=total_start, end_row=header_depth, end_column=total_start)


def _same_header_prefix(left: list[Any], right: list[Any], level: int) -> bool:
    return all(_excel_header_value(left[index]) == _excel_header_value(right[index]) for index in range(level + 1))


def _excel_header_value(value: Any) -> str:
    serialised_value = _serialise_value(value)
    if serialised_value is None:
        return "(Blank)"
    text = str(serialised_value).strip()
    return text if text else "(Blank)"


def _style_header_cell(cell: Any) -> None:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)


def _style_total_cell(cell: Any) -> None:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(fill_type="solid", fgColor=TOTAL_FILL)


def _is_total_header(header: str) -> bool:
    return header == "Total" or header.startswith("Total ") or header.startswith("Grand ")


def _build_export_filename(project_name: str) -> str:
    from datetime import datetime

    safe_project_name = re.sub(r"[^A-Za-z0-9_-]+", "_", project_name).strip("_") or "Project"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"pivot_{safe_project_name}_{timestamp}.xlsx"


def _serialise_value(value: Any) -> str | int | float | bool | None:
    if value is None:
        return None
    if hasattr(value, "item"):
        value = value.item()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, float):
        return round(value, 4)
    return value
